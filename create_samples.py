import os
from reportlab.pdfgen import canvas
from docx import Document

def create_sample_pdf(filename="sample_resumes/sample.pdf"):
    c = canvas.Canvas(filename)
    c.drawString(100, 750, "Name: John Doe")
    c.drawString(100, 730, "Email: john.doe@example.com")
    c.drawString(100, 700, "Experience:")
    c.drawString(100, 680, "- Software Engineer at Tech Corp")
    c.drawString(100, 660, "- Developed Python applications")
    c.drawString(100, 640, "Skills:")
    c.drawString(100, 620, "Python, Machine Learning, Docker, Kubernetes")
    c.save()
    print(f"Created {filename}")

def create_sample_docx(filename="sample_resumes/sample.docx"):
    doc = Document()
    doc.add_heading('Jane Smith', 0)
    doc.add_paragraph('Email: jane.smith@test.com')
    doc.add_heading('Skills', level=1)
    doc.add_paragraph('Java, Spring Boot, React, AWS')
    doc.add_heading('Experience', level=1)
    doc.add_paragraph('Senior Developer at Web Solutions.')
    doc.save(filename)
    print(f"Created {filename}")

if __name__ == "__main__":
    os.makedirs("sample_resumes", exist_ok=True)
    try:
        create_sample_pdf()
        create_sample_docx()
    except ImportError:
        print("Please install 'reportlab' and 'python-docx' to generate samples.")
