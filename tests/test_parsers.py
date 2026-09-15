import os
import sys
import tempfile
import unittest

from docx import Document
from reportlab.pdfgen import canvas

sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from resume_parser.pdf_parser import PDFParser
from resume_parser.word_parser import WordParser


class TestParsers(unittest.TestCase):
    def test_pdf_parser(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            pdf_path = os.path.join(tmpdir, "sample.pdf")
            pdf = canvas.Canvas(pdf_path)
            pdf.drawString(72, 720, "John Doe")
            pdf.drawString(72, 700, "john.doe@example.com")
            pdf.save()

            parser = PDFParser()
            text = parser.parse_file(pdf_path)
            self.assertIn("John Doe", text)
            self.assertIn("john.doe@example.com", text)

    def test_word_parser(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            docx_path = os.path.join(tmpdir, "sample.docx")
            document = Document()
            document.add_paragraph("Jane Smith")
            document.add_paragraph("jane.smith@example.com")
            document.save(docx_path)

            parser = WordParser()
            text = parser.parse_file(docx_path)
            self.assertIn("Jane Smith", text)
            self.assertIn("jane.smith@example.com", text)
